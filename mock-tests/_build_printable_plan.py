#!/usr/bin/env python3
"""Render the 70-day plan as a printable Word document and PDF.

STUDY_PLAN.md is written for reading on screen: one small table per day, 69 of
them, with the section spelled out in full on every row. Printed as-is that runs
to a great many sparse pages. This renders the same plan for paper instead:

  * landscape A4 with narrow margins
  * one continuous table per phase, with a Day column, so the rows flow down the
    page rather than restarting every three tests
  * the section as a single letter against a legend, which frees the width the
    topic column actually needs
  * the header row repeated automatically at the top of every page
  * a blank Done box on each row to tick off as tests are attempted

Data comes from mock-tests/manifest.json, which is the machine-readable source of
truth, while the prose blocks are lifted out of STUDY_PLAN.md so the two cannot
drift apart.

    python3 mock-tests/_build_printable_plan.py

Writes dist/JKP-Constable-Study-Plan.docx and dist/JKP-Constable-Study-Plan.pdf.
"""
import json
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (LongTable, PageBreak, Paragraph, SimpleDocTemplate,
                                Spacer, TableStyle)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "dist")

HEAD = ["Day", "Test #", "Session", "Sec", "Topic / Focus", "Q", "Difficulty", "Type", "Done"]
# millimetres; totals 261 of the 271 usable in landscape A4 at 12.7 mm margins
WIDTHS = [13, 17, 22, 11, 108, 9, 24, 42, 15]

PHASE_TITLES = {
    1: "Phase 1 — Foundation (Topic-wise Building) · Days 1–35",
    2: "Phase 2 — Consolidation (Full-section & First Full-length Mocks) · Days 36–55",
    3: "Phase 3 — Simulation (Daily Full-length Mocks) · Days 56–69",
}


def load():
    m = json.load(open(os.path.join(ROOT, "mock-tests", "manifest.json")))
    tests = sorted(m["tests"], key=lambda t: t["number"])
    plan = open(os.path.join(ROOT, "STUDY_PLAN.md")).read()

    def block(pattern, flags=re.S):
        mm_ = re.search(pattern, plan, flags)
        return mm_.group(1).strip() if mm_ else ""

    notes = {}
    for ph, anchor in ((1, "PHASE 1"), (2, "PHASE 2"), (3, "PHASE 3")):
        notes[ph] = block(rf"## {anchor}[^\n]*\n\n(.+?)\n\n### Day")
    exam = block(r"## Day 70 — EXAM DAY\s*\n\n(.+?)\n\n---")
    tips = re.findall(r"^\d+\.\s+(.+)$",
                      block(r"## Key Tips for Success\s*\n\n(.+?)\n\n\*\*ALL THE BEST"), re.M)
    return tests, notes, exam, tips


def strip_md(s):
    """Markdown emphasis and code ticks do not belong in a printed table."""
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    return s.replace("`", "")


def rows_for(tests, phase):
    out = []
    for t in tests:
        if t.get("phase") != phase:
            continue
        out.append([
            str(t["day"]),
            "%03d" % t["number"],
            t["session"],
            t["subject"],
            t["topics"],
            str(t["total_questions"]),
            t["difficulty_profile"],
            t["type"],
            "",
        ])
    return out


INTRO = [
    ("Paper", "100 MCQs · 1 mark each · 100 marks · 120 minutes · no negative marking prescribed."),
    ("Blueprint", "A General English 25 · B GK & Current Affairs (India) 25 · "
                  "C GK — J&K 10 · D Numerical and Reasoning Ability 25 · "
                  "E Basic Concepts of Computers 15."),
    ("Daily routine", "Revise the topic(s) listed for the day, then attempt that day's tests in "
                      "order: Morning, Afternoon, Late."),
    ("Finding the test file", "The Test # column is the file's name prefix. Test 046 is the file "
                             "046_test_A_...json. Only one file begins with a given three-digit "
                             "number, so scroll to that number and you have the right test."),
    ("Difficulty", "foundation = easy-tilted · standard = blueprint · advanced = hard-tilted · "
                   "exam = full-length simulation. Every section is pitched at 10+2 level — "
                   "difficulty means more steps and subtler distractors, never higher theory."),
]

LEGEND = ("Sec column — A General English · B GK & Current Affairs (India) · C GK with special "
          "reference to J&K · D Numerical and Reasoning Ability · E Basic Concepts of Computers · "
          "FULL all sections together")


# --------------------------------------------------------------------- Word
def repeat_header(row):
    tr = row._tr
    pr = tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def build_docx(tests, notes, exam, tips, path):
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, attr, Mm(12.7))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)

    h = doc.add_heading("70-Day Study & Mock-Test Plan", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("J&K Police Constable (Executive / Armed / IRP / SDRF)  ·  "
                            "207 mock tests  ·  10,350 questions  ·  Exam on Day 70")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(10)
    sub.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    for label, text in INTRO:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label + ": ")
        r.bold = True
        r.font.size = Pt(9)
        p.add_run(text).font.size = Pt(9)

    lg = doc.add_paragraph(LEGEND)
    lg.runs[0].font.size = Pt(8)
    lg.runs[0].italic = True

    for phase in (1, 2, 3):
        doc.add_heading(PHASE_TITLES[phase], level=1)
        if notes.get(phase):
            n = doc.add_paragraph(strip_md(notes[phase]))
            n.runs[0].font.size = Pt(8.5)
            n.runs[0].italic = True

        data = rows_for(tests, phase)
        table = doc.add_table(rows=1, cols=len(HEAD))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        for i, name in enumerate(HEAD):
            c = table.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(name)
            run.bold = True
            run.font.size = Pt(8.5)
            shade(c, "D9E2F3")
        repeat_header(table.rows[0])

        for row in data:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(val)
                run.font.size = Pt(8)
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)

        # widths must be set per cell for Word to honour them
        for r_ in table.rows:
            for i, w in enumerate(WIDTHS):
                r_.cells[i].width = Mm(w)

        if phase != 3:
            doc.add_page_break()

    doc.add_heading("Day 70 — Exam Day", level=1)
    if exam:
        doc.add_paragraph(strip_md(exam)).runs[0].font.size = Pt(9)

    doc.add_heading("Key Tips", level=1)
    for t in tips:
        p = doc.add_paragraph(strip_md(t), style="List Number")
        p.runs[0].font.size = Pt(9)
        p.paragraph_format.space_after = Pt(2)

    doc.save(path)


# ---------------------------------------------------------------------- PDF
def build_pdf(tests, notes, exam, tips, path):
    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], fontSize=18, spaceAfter=2)
    sub = ParagraphStyle("s", parent=ss["Normal"], fontSize=9.5, alignment=1,
                         textColor=colors.HexColor("#444444"), spaceAfter=8)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=8.5, leading=11,
                          alignment=TA_LEFT, spaceAfter=2)
    small = ParagraphStyle("sm", parent=body, fontSize=7.5, textColor=colors.HexColor("#333333"))
    ph = ParagraphStyle("p", parent=ss["Heading2"], fontSize=12, spaceBefore=8, spaceAfter=4)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=7.2, leading=8.6)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold", fontSize=7.4)

    doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                            leftMargin=12.7 * mm, rightMargin=12.7 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title="JKP Constable 70-Day Study Plan")

    flow = [Paragraph("70-Day Study &amp; Mock-Test Plan", title),
            Paragraph("J&amp;K Police Constable (Executive / Armed / IRP / SDRF) &nbsp;·&nbsp; "
                      "207 mock tests &nbsp;·&nbsp; 10,350 questions &nbsp;·&nbsp; Exam on Day 70", sub)]
    for label, text in INTRO:
        flow.append(Paragraph("<b>%s:</b> %s" % (label, text.replace("&", "&amp;")), body))
    flow.append(Paragraph("<i>%s</i>" % LEGEND.replace("&", "&amp;"), small))
    flow.append(Spacer(1, 4))

    style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA5B1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 1.6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")]),
    ])

    for phase in (1, 2, 3):
        flow.append(Paragraph(PHASE_TITLES[phase].replace("&", "&amp;"), ph))
        if notes.get(phase):
            flow.append(Paragraph("<i>%s</i>" % strip_md(notes[phase]).replace("&", "&amp;"), small))
            flow.append(Spacer(1, 3))
        data = [[Paragraph(x, cellb) for x in HEAD]]
        for r in rows_for(tests, phase):
            data.append([Paragraph(str(v).replace("&", "&amp;"), cell) for v in r])
        t = LongTable(data, colWidths=[w * mm for w in WIDTHS], repeatRows=1)
        t.setStyle(style)
        flow.append(t)
        if phase != 3:
            flow.append(PageBreak())

    flow.append(Paragraph("Day 70 — Exam Day", ph))
    if exam:
        flow.append(Paragraph(strip_md(exam).replace("&", "&amp;"), body))
    flow.append(Paragraph("Key Tips", ph))
    for i, tip in enumerate(tips, 1):
        flow.append(Paragraph("%d. %s" % (i, strip_md(tip).replace("&", "&amp;")), body))

    doc.build(flow)


def main():
    tests, notes, exam, tips = load()
    os.makedirs(DIST, exist_ok=True)
    counted = sum(len(rows_for(tests, p)) for p in (1, 2, 3))
    assert counted == len(tests), \
        "%d of %d tests carry no phase and would be dropped" % (len(tests) - counted, len(tests))

    dx = os.path.join(DIST, "JKP-Constable-Study-Plan.docx")
    pf = os.path.join(DIST, "JKP-Constable-Study-Plan.pdf")
    build_docx(tests, notes, exam, tips, dx)
    build_pdf(tests, notes, exam, tips, pf)

    for p in (dx, pf):
        print("wrote %-46s %6.1f KB" % (os.path.relpath(p, ROOT), os.path.getsize(p) / 1024))
    print("rows: %d  (phase 1 %d, phase 2 %d, phase 3 %d)  tips: %d"
          % (counted, len(rows_for(tests, 1)), len(rows_for(tests, 2)),
             len(rows_for(tests, 3)), len(tips)))


if __name__ == "__main__":
    main()
